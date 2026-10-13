import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

YEL = PatternFill('solid', fgColor='FFF7D6'); HEAD = PatternFill('solid', fgColor='1F3A5F'); GREY = PatternFill('solid', fgColor='EEF1F5'); OUT = PatternFill('solid', fgColor='E3F1E7')
BLUE = Font(color='1F4FBF'); BOLD = Font(bold=True); WHITE = Font(bold=True, color='FFFFFF'); thin = Side(style='thin', color='C9CFD8')
INR = '#,##0'; INR2 = '#,##0.00'

class Book:
    def __init__(self, title): self.wb = Workbook(); self.wb.remove(self.wb.active); self.title = title; self.names = {}
    def sheet(self, name, widths):
        ws = self.wb.create_sheet(name)
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
        ws.sheet_view.showGridLines = False; return ws
    def name(self, nm, ws, ref): self.wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ws.title}'!{ref}")
def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t); c.font = WHITE; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center' if i else 'left')
    ws.row_dimensions[row].height = 32
def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = Font(bold=True, size=15, color='1F3A5F')
    if sub: ws['A2'] = sub; ws['A2'].font = Font(italic=True, color='6B7482')

def build(cls, cfg, path):
    ht = cls == 'HT'; B = Book(f'{cls} switchgear')
    # ------------------------------------------------ Read Me
    rm = B.sheet('Read Me', [26, 100]); title(rm, f'{cls} switchgear board: should-cost model', 'Bottom-up cost of a switchboard from its bay schedule. Not a vendor price list: every rate is an editable assumption.')
    lines = [('Purpose', f'Estimate the ex-works price of a {"3.3-33 kV metal-clad VCB board" if ht else "415 V LT board (PCC / MCC / APFC)"} from the schedule of bays. Each bay is costed from its sheet steel, switching device, instrument transformers, protection, wiring and test; the board adds the main busbar{", DC system" if ht else ""}, labour, test, packing, overheads and margin.'),
        ('How to use', 'Edit the yellow cells on Inputs (board rating and the bay schedule) and, if you have quotes, on Rates. Everything else is a live formula. Results are on Summary; the per-bay build-up is on Bay Costing.'),
        ('Sheets', 'Inputs: board specification and bay schedule.  Rates: every price, constant and multiplier (named cells, prefix r_).  Bay Costing: cost per bay type and board-level items.  Summary: cost stack, per-bay and per-ampere figures.  Presets: three sample boards to copy into Inputs.'),
        ('Conventions', 'Yellow fill = input.  Blue font = a typed assumption you may change.  Black = formula.  Green fill = key output.  All amounts are Indian rupees, ex-GST, ex-works.'),
        ('Basis of the numbers', 'Device prices scale with rating from an assumed base price: ' + ('VCB = base(kV) x (A/1250)^0.45 x (kA/25)^0.5. Panel steel, CT, VT and surge-arrester prices are tabulated by voltage class.' if ht else 'ACB = 2.1 lakh x (A/1000)^0.9 x (kA/50)^0.3 (x1.25 for 4-pole); MCCB = 80 x A^1.05 x (kA/36)^0.35; starters, VFDs and APFC banks scale with kW / kVAr.') + ' Steel, copper and aluminium follow commodity prices you type into Rates.'),
        ('Reasonableness', ('About Rs 8-10 lakh per 11 kV panel and Rs 20-25 lakh per 33 kV panel at the default rates.' if ht else 'About Rs 1-4 lakh per LT bay at the default rates (an MCC starter module Rs 0.5-1.5 lakh, an ACB bay Rs 5-12 lakh).') + ' These are market ranges, not quotations. Calibrate Rates with two or three recent offers before relying on a figure.'),
        ('Limits', 'Single board per workbook (copy the file for another). Busbar sized by current density only (no short-circuit or temperature-rise check). No GST, freight to site, erection or civil work. Fault levels above ' + ('40 kA' if ht else '65 kA') + ' need manufacturer confirmation.'),
        ('Companion', 'The same model runs interactively in the Should-Cost dashboards (switchgear.html), with a 3D view and a month selector for the commodity prices.')]
    for i, (a, b) in enumerate(lines, 4):
        rm.cell(i, 1, a).font = BOLD; c = rm.cell(i, 2, b); c.alignment = Alignment(wrap_text=True, vertical='top'); rm.row_dimensions[i].height = 62
    # ------------------------------------------------ Rates
    rt = B.sheet('Rates', [46, 16, 14, 60]); title(rt, 'Rates and constants', 'Blue = assumption. Replace with quotes. Each value is a named cell (r_...), so formulas stay readable.')
    row = [4]
    def sec(t):
        row[0] += 1; c = rt.cell(row[0], 1, t); c.font = BOLD; c.fill = GREY
        for k in (2, 3, 4): rt.cell(row[0], k).fill = GREY
        row[0] += 1
    def P(nm, label, val, unit='', note='', fmt=None):
        r = row[0]; rt.cell(r, 1, label); c = rt.cell(r, 2, val); c.font = BLUE; c.fill = YEL
        if fmt: c.number_format = fmt
        rt.cell(r, 3, unit); rt.cell(r, 4, note).font = Font(color='6B7482', size=9); B.name('r_' + nm, rt, f'$B${r}'); row[0] += 1
    sec('Commodity rates')
    P('copper', 'Copper (busbar)', 1400, 'Rs/kg', 'Commodity file / LME x USDINR'); 
    if not ht: P('aluminium', 'Aluminium (busbar)', 349, 'Rs/kg')
    P('crca', 'CRCA sheet', 76, 'Rs/kg')
    if ht:
        sec('Voltage-class tables')
        r0 = row[0]; hdr(rt, r0, ['kV', 'VCB base Rs (1250 A, 25 kA)', 'Panel steel kg', 'CT set Rs'], 1); rt.cell(r0, 5, 'VT set Rs'); rt.cell(r0, 6, 'Surge arrester Rs'); 
        for k in (5, 6): c = rt.cell(r0, k); c.font = WHITE; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, horizontal='center')
        rt.column_dimensions['E'].width = 14; rt.column_dimensions['F'].width = 14
        tab = [(3.3, 380000, 620, 26000, 60000, 14000), (6.6, 390000, 640, 28000, 66000, 16000), (11, 450000, 700, 34000, 78000, 21000), (22, 720000, 1050, 62000, 140000, 48000), (33, 950000, 1450, 88000, 210000, 70000)]
        for i, t in enumerate(tab, r0 + 1):
            for j, v in enumerate(t, 1):
                c = rt.cell(i, j, v); c.font = BLUE; c.fill = YEL; c.number_format = INR if j > 1 else '0.0'
        e = r0 + len(tab)
        for nm, col in [('kvList', 'A'), ('vcbBase', 'B'), ('panelKg', 'C'), ('ctSet', 'D'), ('vtSet', 'E'), ('laSet', 'F')]: B.name('r_' + nm, rt, f'${col}${r0+1}:${col}${e}')
        row[0] = e + 1
        sec('Switching devices and protection')
        P('vcbAmpExp', 'VCB rating exponent (A/1250)', 0.45); P('vcbKaExp', 'VCB fault-level exponent (kA/25)', 0.5); P('mtrDisc', 'Motor-feeder VCB factor', 0.9, 'x')
        P('ctKaGrow', 'CT price growth per unit of (kA/25 - 1)', 0.15); P('meterHt', 'Multi-function meter', 22000, 'Rs'); P('htAux', 'Earth switch, interlocks, trip-circuit supervision, heaters, lamps', 38000, 'Rs/bay'); P('htAux22', 'Factor for 22 kV and above', 1.6, 'x'); P('ccMult', 'Incomer / coupler relay factor', 1.25, 'x'); P('ptRelayF', 'PT-panel relay factor', 0.7, 'x'); P('vtIncF', 'Incomer VT factor (of VT set)', 0.8, 'x')
        sec('Protection relay prices')
        r0 = row[0]
        for i, (n, v) in enumerate([('Electromechanical', 14000), ('Numeric O/C + E/F', 38000), ('IED, IEC 61850 + disturbance record', 98000)]):
            rt.cell(r0 + i, 1, n); c = rt.cell(r0 + i, 2, v); c.font = BLUE; c.fill = YEL; c.number_format = INR; rt.cell(r0 + i, 3, 'Rs')
        B.name('r_relayList', rt, f'$A${r0}:$A${r0+2}'); B.name('r_relayRate', rt, f'$B${r0}:$B${r0+2}'); row[0] = r0 + 3
        sec('Busbar and DC system')
        P('htJ', 'Busbar current density (copper)', 1.6, 'A/mm2'); P('bayW', 'Panel width', 0.8, 'm'); P('busFab', 'Busbar fabrication, plating', 95, 'Rs/kg'); P('sleeve', 'Sleeving, supports, barriers', 220, 'Rs/m per phase'); P('sleeveHt', 'HT sleeving factor', 1.8, 'x')
        P('dcSystem', '110 V DC system (battery, charger), base', 220000, 'Rs'); P('dcPerBay', 'DC distribution per bay', 14000, 'Rs/bay'); P('arcBoard', 'Arc-flash detection, base', 85000, 'Rs'); P('arcBay', 'Arc-flash detection per bay', 9000, 'Rs/bay'); P('wiringF', 'HT wiring factor', 1.6, 'x'); P('wiring', 'Control wiring and terminals', 6500, 'Rs/bay')
        P('assemblyF', 'HT assembly labour factor', 1.5, 'x'); P('testF', 'HT routine-test factor', 1.8, 'x')
    else:
        sec('Switching devices')
        P('acbBase', 'ACB base price at 1,000 A, 50 kA, 3-pole', 210000, 'Rs'); P('acbAmpExp', 'ACB rating exponent', 0.9); P('acbKaExp', 'ACB fault-level exponent (kA/50)', 0.3); P('acbPoles4', 'Factor for 4-pole (incomers, couplers)', 1.25, 'x')
        P('mccbBase', 'MCCB price = base x A^exp x (kA/36)^exp', 80, 'Rs'); P('mccbAmpExp', 'MCCB rating exponent', 1.05); P('mccbKaExp', 'MCCB fault-level exponent', 0.35)
        P('dolBase', 'DOL starter base (MPCB, contactor, overload)', 3000, 'Rs'); P('dolPerKw', 'DOL starter per kW^exp', 1200, 'Rs'); P('kwExp', 'Starter kW exponent', 0.75); P('sdFactor', 'Star-delta factor', 1.55, 'x'); P('vfdBase', 'VFD = base x kW^exp', 11000, 'Rs'); P('vfdExp', 'VFD kW exponent', 0.8)
        P('apfcBase', 'APFC bank base (controller, switching)', 62000, 'Rs'); P('apfcPerKvar', 'APFC capacitors and reactors per kVAr', 650, 'Rs/kVAr')
        sec('Protection, metering and accessories')
        P('ctLt', 'CT (each)', 4800, 'Rs'); P('relayLt', 'Microprocessor release / relay element', 9500, 'Rs'); P('mfm', 'Multi-function meter', 7500, 'Rs'); P('ammeterSet', 'Ammeter / voltmeter set', 5500, 'Rs'); P('lamps', 'Indicating lamp', 1800, 'Rs'); P('selector', 'Selector switch', 1200, 'Rs'); P('mccbAux', 'MCCB shunt trip / aux', 6000, 'Rs'); P('acbAux', 'ACB shunt trip, closing coil, interlocks', 14000, 'Rs'); P('apfcCtl', 'APFC controller, CTs and metering', 28000, 'Rs')
        sec('Sheet steel per bay (kg)')
        P('acbKgBase', 'ACB bay: base', 160, 'kg'); P('acbKgPerA', 'ACB bay: per ampere', 0.05, 'kg/A'); P('mccbKgBase', 'MCCB feeder: base', 28, 'kg'); P('mccbKgPerA', 'MCCB feeder: per ampere', 0.04, 'kg/A'); P('mccKgBase', 'MCC module: base', 24, 'kg'); P('mccKgPerKw', 'MCC module: per kW^0.8', 0.6, 'kg'); P('vfdKgF', 'VFD module factor', 1.35, 'x'); P('apfcKgBase', 'APFC bay: base', 120, 'kg'); P('apfcKgPerKvar', 'APFC bay: per kVAr', 0.45, 'kg')
        sec('Busbar')
        P('cuA', 'Copper busbar section', 0.8, 'mm2/A'); P('alA', 'Aluminium busbar section', 1.25, 'mm2/A'); P('vertFrac', 'MCC vertical busbar fraction of main', 0.5); P('bayW', 'Column width', 0.8, 'm'); P('busFab', 'Busbar fabrication, plating', 95, 'Rs/kg'); P('sleeve', 'Sleeving, supports, barriers', 220, 'Rs/m per phase'); P('sleeve4b', 'Form 4b sleeving factor', 1.4, 'x'); P('mccPerCol', 'MCC modules per column', 8, 'nos'); P('mccbPerCol', 'MCCB feeders per column', 4, 'nos')
        P('wiring', 'Control wiring and terminals', 6500, 'Rs/bay'); P('mccbWiringF', 'MCCB feeder wiring factor', 0.3, 'x'); P('arcBoard', 'Arc-flash relay and sensors (LT = 50%)', 85000, 'Rs'); P('assemblyF4b', 'Form 4b assembly factor', 1.2, 'x')
    sec('Enclosure multipliers')
    r0 = row[0]
    for i, (n, v) in enumerate([('IP42', 1), ('IP54', 1.1), ('IP55', 1.14)]):
        rt.cell(r0 + i, 1, n); c = rt.cell(r0 + i, 2, v); c.font = BLUE; c.fill = YEL
    B.name('r_ipList', rt, f'$A${r0}:$A${r0+2}'); B.name('r_ipMult', rt, f'$B${r0}:$B${r0+2}'); row[0] = r0 + 3
    if not ht:
        r0 = row[0]
        for i, (n, v) in enumerate([('Form 2b', 1), ('Form 3b', 1.12), ('Form 4b', 1.28)]):
            rt.cell(r0 + i, 1, n); c = rt.cell(r0 + i, 2, v); c.font = BLUE; c.fill = YEL
        B.name('r_formList', rt, f'$A${r0}:$A${r0+2}'); B.name('r_formMult', rt, f'$B${r0}:$B${r0+2}'); row[0] = r0 + 3
    P('arcMult', 'Arc-proof (internal arc classified) enclosure', 1.2, 'x')
    sec('Shop floor, test and cost stack')
    P('powder', 'Pre-treatment and powder coating', 14, 'Rs/kg steel'); P('gasket', 'Gaskets, doors, hardware', 6, 'Rs/kg steel'); P('hdwr', 'Hardware as share of sheet price', 0.06, 'x')
    P('assembly', 'Assembly, fabrication and wiring labour', 38, 'Rs/kg steel'); P('testBay', 'Routine tests, FAT and inspection', 6500, 'Rs/bay'); P('typeTest', 'Type-test and certification allowance', 0.02, 'of material', '', '0.0%')
    P('labelEarth', 'Earth bar, nameplates, end covers', 4500, 'Rs/column'); P('packing', 'Packing', 0.025, 'of ex-works', '', '0.0%'); P('transport', 'Transport to site', 0.02, 'of ex-works', '', '0.0%'); P('overheads', 'Overheads', 0.09, 'of direct', '', '0.0%'); P('margin', 'Margin', 0.10, 'of direct + overheads', '', '0.0%')
    # ------------------------------------------------ Inputs
    ip = B.sheet('Inputs', [44, 16, 16, 40]); title(ip, f'{cls} board: specification and bay schedule')
    def I(r, label, val, unit='', nm=None, dv=None):
        ip.cell(r, 1, label); c = ip.cell(r, 2, val); c.fill = YEL; c.font = BLUE; ip.cell(r, 3, unit)
        if nm: B.name('i_' + nm, ip, f'$B${r}')
        if dv: d = DataValidation(type='list', formula1=dv, allow_blank=False); ip.add_data_validation(d); d.add(c)
        return c
    I(4, 'Tag / board name', cfg['tag'])
    if ht: I(5, 'System voltage', cfg['kv'], 'kV', 'kv', '"3.3,6.6,11,22,33"')
    else: I(5, 'System voltage', 0.415, 'kV (fixed)', 'kv'); ip['B5'].fill = PatternFill(); 
    I(6, 'Busbar rating', cfg['busA'], 'A', 'busA'); I(7, 'Fault level ' + ('(3 s)' if ht else '(1 s)'), cfg['kA'], 'kA', 'kA'); I(8, 'Degree of protection', cfg['ip'], '', 'ip', '"IP42,IP54,IP55"'); I(9, 'Arc-fault protection (IAC)', cfg['arc'], '', 'arc', '"Yes,No"')
    if ht: I(10, 'Protection relay', cfg['relay'], '', 'relay', '"Electromechanical,Numeric O/C + E/F,IED, IEC 61850 + disturbance record"'); ip['D10'] = 'pick from list'
    else: I(10, 'Separation form', cfg['form'], '', 'form', '"Form 2b,Form 3b,Form 4b"'); I(11, 'Busbar material', cfg['busMat'], '', 'busMat', '"Aluminium,Copper"')
    kinds = cfg['kinds']; r0 = 14
    ip.cell(r0 - 1, 1, 'Bay schedule').font = BOLD; hdr(ip, r0, ['Bay type', 'Quantity', 'Size', 'Unit'])
    for i, (k, lab, unit, q, a) in enumerate(kinds):
        r = r0 + 1 + i; ip.cell(r, 1, lab)
        c = ip.cell(r, 2, q); c.fill = YEL; c.font = BLUE
        c = ip.cell(r, 3, a if unit else None); 
        if unit: c.fill = YEL; c.font = BLUE
        ip.cell(r, 4, unit)
    ip.cell(r0 + len(kinds) + 2, 1, 'Size: rated current (A) for breakers and feeders, motor kW for starters, kVAr for APFC banks. Quantity 0 removes a bay type.').font = Font(italic=True, color='6B7482')
    B.name('i_first', ip, f'$B${r0+1}')
    # ------------------------------------------------ Bay Costing
    bc = B.sheet('Bay Costing', [42, 9, 9, 10, 13, 13, 13, 13, 13, 14, 14, 15]); title(bc, 'Cost per bay and board-level items', 'Rows follow the bay schedule on Inputs. Rs per bay unless stated.')
    nk = len(kinds); top = 5
    ipr = lambda i: r0 + 1 + i          # row on Inputs
    mult = f"INDEX(r_ipMult,MATCH(i_ip,r_ipList,0))" + ('' if ht else "*INDEX(r_formMult,MATCH(i_form,r_formList,0))") + "*IF(i_arc=\"Yes\",r_arcMult,1)"
    bc['A3'] = 'Enclosure multiplier (IP x form x arc)'; bc['E3'] = '=' + mult; bc['E3'].number_format = '0.000'; B.name('c_mult', bc, '$E$3')
    if ht:
        hdr(bc, 4, ['Bay type', 'Qty', 'Size A', 'Steel kg / bay', 'Sheet steel', 'Powder coat', 'Gaskets, doors', 'VCB', 'CT set', 'Relay and meter', 'VT / arrester', 'Aux + wiring'])
        kvi = lambda tbl: f"INDEX({tbl},MATCH(i_kv,r_kvList,0))"
        for i, (k, lab, unit, q, a) in enumerate(kinds):
            r = top + i; pt = k == 'pt'
            bc.cell(r, 1, lab); bc.cell(r, 2, f"=Inputs!B{ipr(i)}"); bc.cell(r, 3, f"=Inputs!C{ipr(i)}" if unit else 0)
            bc.cell(r, 4, f"={kvi('r_panelKg')}*{0.8 if pt else 1}*IF(C{r}>1600,1.15,1)*c_mult")
            bc.cell(r, 5, f"=D{r}*r_crca"); bc.cell(r, 6, f"=D{r}*r_powder"); bc.cell(r, 7, f"=D{r}*(r_gasket+r_hdwr*r_crca)")
            if not pt:
                f = {'inc': 'r_ccMult', 'cpl': 'r_ccMult'}.get(k)
                vcb = f"={kvi('r_vcbBase')}*(C{r}/1250)^r_vcbAmpExp*(i_kA/25)^r_vcbKaExp" + ('*r_mtrDisc' if k == 'mtr' else '')
                bc.cell(r, 8, vcb); bc.cell(r, 9, f"={kvi('r_ctSet')}*(1+r_ctKaGrow*(i_kA/25-1))")
                bc.cell(r, 10, f"=INDEX(r_relayRate,MATCH(i_relay,r_relayList,0))*{'r_ccMult' if f else 1}+r_meterHt")
                bc.cell(r, 11, f"={kvi('r_vtSet')}*r_vtIncF" if k == 'inc' else 0)
                bc.cell(r, 12, f"=r_htAux*IF(i_kv>=22,r_htAux22,1)+r_wiring*r_wiringF")
            else:
                bc.cell(r, 8, 0); bc.cell(r, 9, 0); bc.cell(r, 10, f"=INDEX(r_relayRate,MATCH(i_relay,r_relayList,0))*r_ptRelayF+r_meterHt")
                bc.cell(r, 11, f"={kvi('r_vtSet')}+{kvi('r_laSet')}"); bc.cell(r, 12, "=r_wiring*r_wiringF")
        end = top + nk - 1; col_t = 'M'; bc.column_dimensions['M'].width = 15; bc.column_dimensions['N'].width = 16
        for c_, t in (('M', 'Total per bay'), ('N', 'Total (qty x)')):
            c = bc[f'{c_}4']; c.value = t; c.font = WHITE; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, horizontal='center')
        for r in range(top, end + 1):
            bc[f'M{r}'] = f"=SUM(E{r}:L{r})"; bc[f'N{r}'] = f"=B{r}*M{r}"
        grp = {'Enclosure & structure': f"SUMPRODUCT(B{top}:B{end},E{top}:E{end}+F{top}:F{end}+G{top}:G{end})", 'Switching devices': f"SUMPRODUCT(B{top}:B{end},H{top}:H{end})", 'Protection & metering': f"SUMPRODUCT(B{top}:B{end},I{top}:I{end}+J{top}:J{end}+K{top}:K{end})", 'Wiring & auxiliaries (bays)': f"SUMPRODUCT(B{top}:B{end},L{top}:L{end})"}
        steelTot = f"SUMPRODUCT(B{top}:B{end},D{top}:D{end})"; nBays = f"SUM(B{top}:B{end})"; cols = nBays
    else:
        hdr(bc, 4, ['Bay type', 'Qty', 'Size', 'Steel kg / bay', 'Sheet steel', 'Powder coat', 'Gaskets, doors', 'Device', 'Protection & metering', 'Aux + wiring', 'Total per bay', 'Total (qty x)'])
        for i, (k, lab, unit, q, a) in enumerate(kinds):
            r = top + i
            bc.cell(r, 1, lab); bc.cell(r, 2, f"=Inputs!B{ipr(i)}"); bc.cell(r, 3, f"=Inputs!C{ipr(i)}")
            S = f"C{r}"
            if k in ('aci', 'acc', 'acf'):
                steel = f"=(r_acbKgBase+r_acbKgPerA*{S})*c_mult"; dev = f"=r_acbBase*({S}/1000)^r_acbAmpExp*(i_kA/50)^r_acbKaExp*{1 if k == 'acf' else 'r_acbPoles4'}"
                prot = "=r_ctLt*4+r_relayLt*3+r_mfm+r_ammeterSet"; aux = "=r_acbAux+r_lamps*3+r_selector*2+r_wiring"
            elif k == 'mcf':
                steel = f"=(r_mccbKgBase+r_mccbKgPerA*{S})*c_mult"; dev = f"=r_mccbBase*{S}^r_mccbAmpExp*(i_kA/36)^r_mccbKaExp"; prot = "=r_ctLt+r_mccbAux+r_lamps"; aux = "=r_wiring*r_mccbWiringF"
            elif k in ('dol', 'sd', 'vfd'):
                steel = f"=(r_mccKgBase+r_mccKgPerKw*{S}^0.8)*c_mult*{'r_vfdKgF' if k == 'vfd' else 1}"
                dev = f"=r_vfdBase*{S}^r_vfdExp" if k == 'vfd' else f"=(r_dolBase+r_dolPerKw*{S}^r_kwExp)*{'r_sdFactor' if k == 'sd' else 1}"
                prot = "=r_ctLt+r_relayLt+r_lamps*3+r_selector*2"; aux = "=r_wiring"
            else:
                steel = f"=(r_apfcKgBase+r_apfcKgPerKvar*{S})*c_mult"; dev = f"=r_apfcBase+r_apfcPerKvar*{S}"; prot = "=r_apfcCtl"; aux = "=r_wiring"
            bc.cell(r, 4, steel); bc.cell(r, 5, f"=D{r}*r_crca"); bc.cell(r, 6, f"=D{r}*r_powder"); bc.cell(r, 7, f"=D{r}*(r_gasket+r_hdwr*r_crca)"); bc.cell(r, 8, dev); bc.cell(r, 9, prot); bc.cell(r, 10, aux)
            bc.cell(r, 11, f"=SUM(E{r}:J{r})"); bc.cell(r, 12, f"=B{r}*K{r}")
            bc.cell(r, 13, f"=B{r}/r_mccbPerCol" if k == 'mcf' else (f"=B{r}/r_mccPerCol" if k in ('dol', 'sd', 'vfd') else f"=B{r}"))
        end = top + nk - 1
        grp = {'Enclosure & structure': f"SUMPRODUCT(B{top}:B{end},E{top}:E{end}+F{top}:F{end}+G{top}:G{end})", 'Switching devices': f"SUMPRODUCT(B{top}:B{end},H{top}:H{end})", 'Protection & metering': f"SUMPRODUCT(B{top}:B{end},I{top}:I{end})", 'Wiring & auxiliaries (bays)': f"SUMPRODUCT(B{top}:B{end},J{top}:J{end})"}
        steelTot = f"SUMPRODUCT(B{top}:B{end},D{top}:D{end})"; nBays = f"SUM(B{top}:B{end})"
        bc.column_dimensions['M'].width = 10; c = bc['M4']; c.value = 'Columns'; c.font = WHITE; c.fill = HEAD; c.alignment = Alignment(horizontal='center')
        cols = f"SUM(M{top}:M{end})"
    # board-level block
    b0 = end + 3; bc.cell(b0 - 1, 1, 'Board-level items').font = BOLD; hdr(bc, b0, ['Item', 'Qty', 'Unit', '', 'Rate Rs', '', '', '', 'Amount Rs'] if ht else ['Item', 'Qty', 'Unit', '', 'Rate Rs', '', '', '', 'Amount Rs'])
    rows = []
    def BL(label, qty, unit, rate, group):
        r = b0 + 1 + len(rows); bc.cell(r, 1, label); bc.cell(r, 2, '=' + qty); bc.cell(r, 3, unit); bc.cell(r, 5, '=' + rate); bc.cell(r, 9, f"=B{r}*E{r}"); rows.append((r, group, label))
        bc.cell(r, 10, group).font = Font(color='6B7482', size=9)
        for k in (2, 5, 9): bc.cell(r, k).number_format = INR2 if k == 2 else INR
        return r
    bc.cell(b0, 10, 'Group').font = WHITE; bc.cell(b0, 10).fill = HEAD
    busMm2 = "i_busA/r_htJ" if ht else "i_busA*IF(i_busMat=\"Copper\",r_cuA,r_alA)"
    ph = 3 if ht else 4; dens = "8.9" if ht else "IF(i_busMat=\"Copper\",8.9,2.7)"; raw = "r_copper" if ht else "IF(i_busMat=\"Copper\",r_copper,r_aluminium)"
    bc['A' + str(b0 - 2)] = None
    # helper cells
    h0 = b0 + 14
    bc.cell(h0, 1, 'Helper figures').font = BOLD
    bc.cell(h0 + 1, 1, 'Number of bays (units)'); bc.cell(h0 + 1, 2, '=' + nBays)
    bc.cell(h0 + 2, 1, 'Columns / panels'); bc.cell(h0 + 2, 2, '=' + (nBays if ht else cols))
    bc.cell(h0 + 3, 1, 'Board length m (columns x width)'); bc.cell(h0 + 3, 2, f"=B{h0+2}*r_bayW")
    bc.cell(h0 + 4, 1, 'Busbar run length m (board + 1.2 m risers per column)'); bc.cell(h0 + 4, 2, f"=B{h0+3}+1.2*ROUNDUP(B{h0+2},0)")
    bc.cell(h0 + 5, 1, 'Busbar section mm2 per phase'); bc.cell(h0 + 5, 2, '=' + busMm2)
    bc.cell(h0 + 6, 1, 'Sheet steel total kg'); bc.cell(h0 + 6, 2, '=' + steelTot)
    for k in range(1, 7): bc.cell(h0 + k, 2).number_format = INR2
    nb, ncol, blen, bsec, steelC = f"B{h0+1}", f"B{h0+2}", f"B{h0+4}", f"B{h0+5}", f"B{h0+6}"
    BL(f"Main busbar ({'copper' if ht else 'per busbar material'})", f"{ph}*{bsec}*{blen}*{dens}/1000", 'kg', f"{raw}+r_busFab", 'Busbars & insulation')
    if not ht:
        mccRows = [top + i for i, (k, *_ ) in enumerate(kinds) if k in ('dol', 'sd', 'vfd')]
        BL('Vertical busbars, MCC sections', f"IF(({'+'.join('B'+str(r) for r in mccRows)})>0,{ph}*{bsec}*r_vertFrac*2.2*ROUNDUP(({'+'.join('B'+str(r) for r in mccRows)})/r_mccPerCol,0)*{dens}/1000,0)", 'kg', f"{raw}+r_busFab", 'Busbars & insulation')
    BL('Busbar sleeving, supports and barriers', f"{blen}*{ph}", 'm', "r_sleeve" + ("*r_sleeveHt" if ht else "*IF(i_form=\"Form 4b\",r_sleeve4b,1)"), 'Busbars & insulation')
    if ht:
        BL('110 V DC system', '1', 'set', f"r_dcSystem+r_dcPerBay*{nb}", 'Wiring & auxiliaries'); BL('Arc-flash detection and fast tripping', 'IF(i_arc="Yes",1,0)', 'set', f"r_arcBoard+r_arcBay*{nb}", 'Wiring & auxiliaries')
    else:
        BL('Arc-flash relay and sensors', 'IF(i_arc="Yes",1,0)', 'set', 'r_arcBoard*0.5', 'Wiring & auxiliaries')
    BL('Earth bar, nameplates, end covers', f"ROUNDUP({ncol},0)", 'column', 'r_labelEarth', 'Wiring & auxiliaries')
    # material subtotal (before G6)
    mat = f"(SUM(B{top}:B{end}*0)+" if False else ''
    bayTot = ('N' if ht else 'L')
    bays_sum = f"SUM({bayTot}{top}:{bayTot}{end})"
    boardPre = '+'.join(f"I{r}" for r, g, l in rows)
    BL('Assembly, fabrication and wiring labour', steelC, 'kg', "r_assembly" + ("*r_assemblyF" if ht else "*IF(i_form=\"Form 4b\",r_assemblyF4b,1)"), 'Assembly & test')
    BL('Routine tests, FAT and inspection', nb, 'bay', "r_testBay" + ("*r_testF" if ht else ""), 'Assembly & test')
    BL('Type-test and certification allowance', '1', 'lot', f"({bays_sum}+{boardPre})*r_typeTest", 'Assembly & test')
    BL('Packing and transport', '1', 'lot', f"({bays_sum}+SUM(I{b0+1}:I{b0+len(rows)}))*(r_packing+r_transport)", 'Assembly & test')
    last = b0 + len(rows)
    # ------------------------------------------------ Summary
    sm = B.sheet('Summary', [44, 20, 14, 46]); title(sm, f'{cls} board: should-cost summary'); sm['A3'] = '=Inputs!B4'; sm['A3'].font = BOLD
    hdr(sm, 5, ['Cost group', 'Rs', 'Share', 'Note'])
    bsum = lambda g: f"SUMIF('Bay Costing'!J{b0+1}:J{last},\"{g}\",'Bay Costing'!I{b0+1}:I{last})"
    rr = 6
    items = [('Enclosure & structure', grp['Enclosure & structure'].replace('B', "'Bay Costing'!B").replace('E', "'Bay Costing'!E").replace('F', "'Bay Costing'!F").replace('G', "'Bay Costing'!G") if False else None)]
    def bsp(expr):  # prefix sheet refs for bay columns
        import re
        return re.sub(r"\b([B-N])(\d+):([B-N])(\d+)", lambda m: f"'Bay Costing'!{m.group(1)}{m.group(2)}:{m.group(3)}{m.group(4)}", expr)
    stack = [('Enclosure & structure', '=' + bsp(grp['Enclosure & structure']), 'Sheet steel, pre-treatment, powder coat, doors'), ('Busbars & insulation', '=' + bsum('Busbars & insulation'), 'Main (and vertical) busbar, sleeving'), ('Switching devices', '=' + bsp(grp['Switching devices']), 'VCBs' if ht else 'ACBs, MCCBs, starters, VFDs, APFC'),
             ('Protection & metering', '=' + bsp(grp['Protection & metering']), 'CTs, VTs, relays, meters' + (', arresters' if ht else '')), ('Wiring & auxiliaries', '=' + bsp(grp['Wiring & auxiliaries (bays)']) + '+' + bsum('Wiring & auxiliaries'), 'Control wiring, interlocks, ' + ('DC system, arc detection' if ht else 'lamps, selectors')), ('Assembly, test, packing & transport', '=' + bsum('Assembly & test'), 'Labour, routine and type tests, packing')]
    for lab, f, note in stack:
        sm.cell(rr, 1, lab); sm.cell(rr, 2, f).number_format = INR; sm.cell(rr, 4, note).font = Font(color='6B7482', size=9); rr += 1
    sm.cell(rr, 1, 'Direct cost').font = BOLD; sm.cell(rr, 2, f"=SUM(B6:B{rr-1})").number_format = INR; sm.cell(rr, 2).font = BOLD; dr = rr; rr += 1
    sm.cell(rr, 1, 'Overheads'); sm.cell(rr, 2, f"=B{dr}*r_overheads").number_format = INR; orr = rr; rr += 1
    sm.cell(rr, 1, 'Margin'); sm.cell(rr, 2, f"=(B{dr}+B{orr})*r_margin").number_format = INR; rr += 1
    sm.cell(rr, 1, 'SHOULD-COST, ex-works (Rs)').font = BOLD; c = sm.cell(rr, 2, f"=B{dr}+B{orr}+B{rr-1}"); c.number_format = INR; c.font = BOLD; c.fill = OUT; tot = rr; rr += 1
    for r in range(6, tot + 1): sm.cell(r, 3, f"=B{r}/B${tot}").number_format = '0.0%'
    sm.cell(rr, 1, 'Should-cost (Rs crore)'); sm.cell(rr, 2, f"=B{tot}/10000000").number_format = '0.000'; rr += 1
    sm.cell(rr, 1, 'Number of bays (units)'); sm.cell(rr, 2, f"='Bay Costing'!B{h0+1}"); rr += 1
    sm.cell(rr, 1, 'Cost per bay (Rs)'); sm.cell(rr, 2, f"=B{tot}/B{rr-1}").number_format = INR; rr += 1
    sm.cell(rr, 1, 'Cost per ampere of busbar (Rs/A)'); sm.cell(rr, 2, f"=B{tot}/Inputs!B6").number_format = INR; rr += 1
    sm.cell(rr, 1, 'Sheet steel (kg)'); sm.cell(rr, 2, f"='Bay Costing'!B{h0+6}").number_format = INR; rr += 1
    sm.cell(rr, 1, 'Busbar mass (kg)'); sm.cell(rr, 2, f"='Bay Costing'!B{b0+1}" + ("" if ht else f"+'Bay Costing'!B{b0+2}")).number_format = INR; rr += 2
    sm.cell(rr, 1, 'Checks').font = BOLD; rr += 1
    sm.cell(rr, 1, 'Largest bay rating within the busbar rating?')
    ratingRows = [i for i, (k, lab, unit, *_ ) in enumerate(kinds) if unit == 'A']
    sm.cell(rr, 2, '=IF(MAX(' + ','.join(f"Inputs!C{ipr(i)}*(Inputs!B{ipr(i)}>0)" for i in ratingRows) + ')<=Inputs!B6,"OK","CHECK: a bay exceeds the busbar")'); rr += 1
    sm.cell(rr, 1, 'Fault level within the usual range?'); sm.cell(rr, 2, f'=IF(Inputs!B7<={40 if ht else 65},"OK","CHECK: confirm with manufacturer")')
    B.name('o_total', sm, f'$B${tot}')
    # ------------------------------------------------ Presets
    ps = B.sheet('Presets', [48, 10, 10, 10, 14, 10, 40]); title(ps, 'Sample boards (copy the values into Inputs)')
    hdr(ps, 4, ['Board', 'kV', 'Busbar A', 'kA'] + [''] * 3)
    for i, p in enumerate(cfg['presets'], 5):
        ps.cell(i, 1, p['tag']); ps.cell(i, 2, p.get('kv', 0.415)); ps.cell(i, 3, p['busA']); ps.cell(i, 4, p['kA']); ps.cell(i, 7, p['note']).font = Font(color='6B7482', size=9)
        ps.cell(i, 5, '; '.join(f"{k}={q}" + (f"@{a}" if a else '') for k, q, a in p['bays'])).alignment = Alignment(wrap_text=True)
        ps.merge_cells(start_row=i, start_column=5, end_row=i, end_column=6); ps.row_dimensions[i].height = 44
    for ws in B.wb.worksheets:
        if ws.title in ('Bay Costing',): ws.freeze_panes = 'B5'
        for row_ in ws.iter_rows():
            for c in row_:
                if c.number_format == 'General' and isinstance(c.value, str) and c.value.startswith('=') and c.column > 1 and ws.title in ('Bay Costing',): c.number_format = INR
    B.wb.save(path); return path
